from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
from docx import Document
from pptx import Presentation
from pypdf import PdfReader


class DocumentExtractionError(ValueError):
    """Raised when a document cannot be converted to plain text."""


SUPPORTED_EXTENSIONS = {
    ".csv",
    ".docx",
    ".md",
    ".pdf",
    ".pptx",
    ".txt",
    ".xlsx",
}


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentExtractionError(f"Unsupported file type '{suffix}'. Supported: {supported}.")

    if suffix in {".csv", ".md", ".txt"}:
        return _extract_plain_text(content)
    if suffix == ".pdf":
        return _extract_pdf(content)
    if suffix == ".docx":
        return _extract_docx(content)
    if suffix == ".pptx":
        return _extract_pptx(content)
    if suffix == ".xlsx":
        return _extract_xlsx(content)

    raise DocumentExtractionError(f"Unsupported file type '{suffix}'.")


def _extract_plain_text(content: bytes) -> str:
    try:
        return content.decode("utf-8").strip()
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError("Plain text files must be UTF-8 encoded.") from exc


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(page.strip() for page in pages if page.strip())


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    return "\n".join(paragraph for paragraph in paragraphs if paragraph)


def _extract_pptx(content: bytes) -> str:
    presentation = Presentation(BytesIO(content))
    parts: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    slide_text.append(text)
        if slide_text:
            parts.append(f"Slide {slide_number}\n" + "\n".join(slide_text))
    return "\n\n".join(parts)


def _extract_xlsx(content: bytes) -> str:
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet_text: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) for value in row if value is not None]
            if values:
                rows.append(" | ".join(values))
        if rows:
            sheet_text.append(f"Sheet: {sheet.title}\n" + "\n".join(rows))
    return "\n\n".join(sheet_text)
